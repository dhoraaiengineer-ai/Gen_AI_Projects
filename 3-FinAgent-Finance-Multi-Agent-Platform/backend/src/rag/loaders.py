"""Document loaders: every source becomes a list of located ``Section`` objects.

Supported: PDF (text + OCR for scanned pages), DOCX, XLSX, CSV/TSV, PPTX, Markdown, TXT, JSON,
XML, HTML, SQL scripts, e-mail (.eml, attachments routed to their loaders) and images (OCR).
Records from REST APIs, SQL query results and web search results load through
``load_records`` / ``load_web_results``. Legacy binary Office formats, encrypted PDFs and files
over the size limit are rejected with a clear message. Text is normalised (CRLF → LF).
"""

from __future__ import annotations

import csv
import email
import io
import json
import logging
import re
from collections.abc import Callable, Iterable, Mapping, Sequence
from email import policy
from html.parser import HTMLParser
from pathlib import PurePath
from typing import Any, ClassVar
from xml.etree import ElementTree

import sqlparse

from src.core.errors import ValidationFailed
from src.rag.types import LoadedDocument, Ocr, Section

logger = logging.getLogger(__name__)

MAX_BYTES = 25 * 1024 * 1024
LEGACY = {".xls": ".xlsx", ".doc": ".docx", ".ppt": ".pptx"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}
_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*$")


def normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    return re.sub(r"[ \t]+\n", "\n", text).strip()


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def table_markdown(rows: Sequence[Sequence[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    padded = [list(r) + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(padded[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in padded[1:]]
    return "\n".join(lines)


def _table_section(rows: list[list[str]], location: dict[str, Any], heading: list[str]) -> Section:
    rows = [r for r in rows if any(c for c in r)]
    return Section(table_markdown(rows), "table", location, heading, rows)


# ------------------------------------------------------------------------------ loaders
def _pdf(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    from pypdf import PdfReader  # noqa: PLC0415

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ValidationFailed(f"{name}: password-protected PDFs are not supported")
    sections: list[Section] = []
    warnings: list[str] = []
    needs_ocr = False
    for i, page in enumerate(reader.pages):
        text = normalise(page.extract_text() or "")
        location: dict[str, Any] = {"page": i + 1}
        if len(text) < 20:
            if ocr.available():
                text = normalise(ocr.pdf_page_text(data, i))
                location["ocr"] = True
            else:
                needs_ocr = True
                warnings.append(f"page {i + 1} has no text layer (OCR unavailable)")
                continue
        first = text.split("\n", 1)[0].strip()
        heading = [first] if 0 < len(first) <= 80 and not first.endswith(".") else []
        sections.append(Section(text, "text", location, heading))
    title = (reader.metadata.title if reader.metadata else None) or PurePath(name).stem
    return LoadedDocument(name, "pdf", str(title), sections, len(reader.pages), needs_ocr, warnings)


def _docx(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    import docx  # noqa: PLC0415
    from docx.table import Table  # noqa: PLC0415
    from docx.text.paragraph import Paragraph  # noqa: PLC0415

    document = docx.Document(io.BytesIO(data))
    sections: list[Section] = []
    path: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            text = normalise("\n".join(buffer))
            if text:
                label = path[-1] if path else "Introduction"
                sections.append(Section(text, "text", {"section": label}, list(path)))
            buffer.clear()

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            style = (block.style.name if block.style is not None else "") or ""
            text = block.text.strip()
            if not text:
                continue
            if style.startswith("Heading") or style == "Title":
                flush()
                level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
                path = [*path[: level - 1], text]
            else:
                buffer.append(text)
        elif isinstance(block, Table):
            flush()
            rows = [[c.text.strip() for c in row.cells] for row in block.rows]
            label = path[-1] if path else "Table"
            sections.append(_table_section(rows, {"section": label}, list(path)))
    flush()
    title = document.core_properties.title or (path[0] if path else PurePath(name).stem)
    return LoadedDocument(name, "docx", title, sections)


def _xlsx(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    import openpyxl  # noqa: PLC0415

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections = []
    try:
        for sheet in workbook.worksheets:
            rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            rows = [r for r in rows if any(r)]
            if rows:
                sections.append(_table_section(rows, {"sheet": sheet.title}, [sheet.title]))
    finally:
        workbook.close()
    return LoadedDocument(name, "xlsx", PurePath(name).stem, sections)


def _csv(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    text = normalise(data.decode("utf-8-sig", errors="replace"))
    dialect = csv.excel_tab if name.lower().endswith(".tsv") else csv.excel
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    return LoadedDocument(name, "csv", PurePath(name).stem, [_table_section(rows, {}, [])])


def _pptx(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    from pptx import Presentation  # noqa: PLC0415

    deck = Presentation(io.BytesIO(data))
    sections: list[Section] = []
    for n, slide in enumerate(deck.slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text.strip() for c in row.cells] for row in shape.table.rows]
                parts.append(table_markdown(rows))
            if shape.shape_type == 13 and ocr.available():  # picture → multimodal chunk
                text = normalise(ocr.image_text(shape.image.blob))
                if text:
                    loc = {"slide": n, "figure": shape.name, "ocr": True}
                    sections.append(Section(text, "image", loc, [title] if title else []))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        body = normalise("\n".join([title, *parts]) if title else "\n".join(parts))
        if body:
            sections.append(Section(body, "slide", {"slide": n}, [title] if title else []))
    deck_title = deck.core_properties.title or PurePath(name).stem
    return LoadedDocument(name, "pptx", deck_title, sections)


def markdown_sections(text: str, location_key: str = "section") -> list[Section]:
    sections: list[Section] = []
    path: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        body = normalise("\n".join(buffer))
        if body:
            label = path[-1] if path else "Introduction"
            sections.append(Section(body, "text", {location_key: label}, list(path)))
        buffer.clear()

    for line in normalise(text).split("\n"):
        match = _MD_HEADING.match(line)
        if match:
            flush()
            level = len(match.group(1))
            path = [*path[: level - 1], match.group(2).strip()]
        else:
            buffer.append(line)
    flush()
    return sections


def _markdown(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    text = data.decode("utf-8-sig", errors="replace")
    sections = markdown_sections(text)
    title = (
        sections[0].heading_path[0]
        if sections and sections[0].heading_path
        else PurePath(name).stem
    )
    return LoadedDocument(name, "md", title, sections)


def _txt(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    text = normalise(data.decode("utf-8-sig", errors="replace"))
    return LoadedDocument(name, "txt", PurePath(name).stem, [Section(text, "text", {}, [])])


def flatten(value: Any, prefix: str = "") -> list[str]:
    """``key: value`` lines for a JSON-like record."""
    if isinstance(value, Mapping):
        return [line for k, v in value.items() for line in flatten(v, f"{prefix}{k}.")]
    if isinstance(value, list) and value and all(not isinstance(v, Mapping | list) for v in value):
        return [f"{prefix.rstrip('.')}: {', '.join(str(v) for v in value)}"]
    if isinstance(value, list):
        return [line for i, v in enumerate(value) for line in flatten(v, f"{prefix}{i}.")]
    return [f"{prefix.rstrip('.')}: {value}"]


def load_records(
    name: str,
    records: Iterable[Mapping[str, Any]],
    *,
    source_type: str = "api",
    key: str | None = None,
) -> LoadedDocument:
    """REST API responses or SQL query results → one ``record`` section per row."""
    sections = []
    for i, record in enumerate(records):
        ident = str(record.get(key)) if key and key in record else str(i)
        loc = {"path": f"{name}[{ident}]"}
        sections.append(Section("\n".join(flatten(record)), "record", loc, []))
    return LoadedDocument(name, source_type, name, sections)


def _json(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    payload = json.loads(data.decode("utf-8-sig"))
    items: list[tuple[str, Any]] = []
    if isinstance(payload, list):
        items = [(f"$[{i}]", v) for i, v in enumerate(payload)]
    elif isinstance(payload, dict):
        for key, value in payload.items():
            if isinstance(value, list) and value and isinstance(value[0], Mapping):
                items += [(f"$.{key}[{i}]", v) for i, v in enumerate(value)]
            else:
                items.append((f"$.{key}", {key: value}))
    sections = [Section("\n".join(flatten(v)), "record", {"path": p}, []) for p, v in items]
    title = (
        payload.get("title", PurePath(name).stem)
        if isinstance(payload, dict)
        else PurePath(name).stem
    )
    return LoadedDocument(name, "json", str(title), sections)


def _xml(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    root = ElementTree.fromstring(data)  # noqa: S314 — expat ≥ 2.4.1 limits entity expansion

    def lines(el: ElementTree.Element, prefix: str) -> list[str]:
        out = [f"{prefix}{el.tag}: {el.text.strip()}"] if el.text and el.text.strip() else []
        out += [f"{prefix}{el.tag}.@{k}: {v}" for k, v in el.attrib.items()]
        for child in el:
            out += lines(child, f"{prefix}{el.tag}.")
        return out

    sections = [
        Section(
            "\n".join(lines(child, "")), "record", {"path": f"/{root.tag}/{child.tag}[{i}]"}, []
        )
        for i, child in enumerate(root)
    ]
    return LoadedDocument(name, "xml", PurePath(name).stem, sections)


class _HtmlText(HTMLParser):
    SKIP: ClassVar[frozenset[str]] = frozenset(
        {"script", "style", "nav", "footer", "header", "aside", "noscript", "form"}
    )
    HEADINGS: ClassVar[dict[str, int]] = {"h1": 1, "h2": 2, "h3": 3}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0
        self.heading: int | None = None
        self.title = ""
        self.in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self.SKIP:
            self.skip += 1
        elif tag in self.HEADINGS:
            self.heading = self.HEADINGS[tag]
            self.out.append("\n" + "#" * self.heading + " ")
        elif tag == "title":
            self.in_title = True
        elif tag in {"p", "br", "li", "tr", "div"}:
            self.out.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        elif tag in self.HEADINGS:
            self.heading = None
            self.out.append("\n")
        elif tag == "title":
            self.in_title = False

    def handle_data(self, data: str) -> None:
        if self.in_title:
            self.title += data
        elif not self.skip:
            self.out.append(data if self.heading is None else data.strip())


def html_to_markdown(html: str) -> tuple[str, str]:
    parser = _HtmlText()
    parser.feed(html)
    return parser.title.strip(), "".join(parser.out)


def _html(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    title, text = html_to_markdown(data.decode("utf-8", errors="replace"))
    return LoadedDocument(name, "html", title or PurePath(name).stem, markdown_sections(text))


def load_web_results(items: Iterable[Mapping[str, Any]]) -> list[LoadedDocument]:
    """Web search results (title, url, content/snippet, published) → one document per URL."""
    docs = []
    for item in items:
        body = str(item.get("content") or item.get("snippet") or "")
        title, text = html_to_markdown(body) if "<" in body else ("", body)
        url = str(item.get("url", "web"))
        sections = markdown_sections(text) or [Section(normalise(text), "text", {}, [])]
        for s in sections:
            s.location = {"url": url, **s.location}
        docs.append(LoadedDocument(url, "web", str(item.get("title") or title or url), sections))
    return docs


def _sql(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    text = normalise(data.decode("utf-8-sig", errors="replace"))
    statements = [s.strip() for s in sqlparse.split(text) if s.strip()]
    sections = [Section(s, "code", {"statement": i}, []) for i, s in enumerate(statements, start=1)]
    return LoadedDocument(name, "sql", PurePath(name).stem, sections)


def _image(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    if not ocr.available():
        return LoadedDocument(
            name, "image", PurePath(name).stem, [], needs_ocr=True, warnings=["OCR unavailable"]
        )
    text = normalise(ocr.image_text(data))
    sections = [Section(text, "image", {"figure": name, "ocr": True}, [])] if text else []
    return LoadedDocument(name, "image", PurePath(name).stem, sections)


def _eml(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    message = email.message_from_bytes(data, policy=policy.default)
    subject = str(message.get("subject", PurePath(name).stem))
    body = message.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        content = body.get_content()
        text = html_to_markdown(content)[1] if body.get_content_type() == "text/html" else content
    header = f"From: {message.get('from', '')}\nDate: {message.get('date', '')}\nSubject: {subject}"
    sections = [
        Section(normalise(f"{header}\n\n{text}"), "text", {"section": "message"}, [subject])
    ]
    warnings: list[str] = []
    for part in message.iter_attachments():
        filename = part.get_filename() or "attachment"
        try:
            payload = part.get_payload(decode=True)
            if isinstance(payload, bytes):
                attached = load_bytes(payload, filename, ocr)
                for s in attached.sections:
                    s.location = {"attachment": filename, **s.location}
                sections += attached.sections
        except ValidationFailed as exc:
            warnings.append(f"{filename}: {exc.message}")
    return LoadedDocument(name, "eml", subject, sections, warnings=warnings)


LOADERS: dict[str, Callable[[bytes, str, Ocr], LoadedDocument]] = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".xlsx": _xlsx,
    ".csv": _csv,
    ".tsv": _csv,
    ".pptx": _pptx,
    ".md": _markdown,
    ".markdown": _markdown,
    ".txt": _txt,
    ".json": _json,
    ".xml": _xml,
    ".html": _html,
    ".htm": _html,
    ".sql": _sql,
    ".eml": _eml,
    **dict.fromkeys(IMAGE_EXT, _image),
}


def load_bytes(data: bytes, name: str, ocr: Ocr) -> LoadedDocument:
    ext = PurePath(name).suffix.lower()
    if ext in LEGACY:
        raise ValidationFailed(
            f"{name}: legacy {ext} files are not supported; save as {LEGACY[ext]}"
        )
    if len(data) > MAX_BYTES:
        raise ValidationFailed(f"{name}: larger than {MAX_BYTES // (1024 * 1024)} MB")
    loader = LOADERS.get(ext)
    if loader is None:
        raise ValidationFailed(f"{name}: unsupported file type {ext or '(none)'}")
    try:
        return loader(data, name, ocr)
    except ValidationFailed:
        raise
    except Exception as exc:
        logger.warning("loader failed", extra={"source": name, "error": type(exc).__name__})
        raise ValidationFailed(f"{name}: could not be read ({type(exc).__name__})") from exc
