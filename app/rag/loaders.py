"""Turn uploaded files into text sections, keeping where each piece came from (page, sheet, slide).

Each section's metadata ends up on its chunks, so citations can say "report.pdf, p. 3".
Tables (Excel sheets, Word/PowerPoint tables) become markdown tables, which the table chunker
splits between rows with the header repeated.
"""

import io
import logging
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from pathlib import PurePath

from app.rag import ocr

logger = logging.getLogger(__name__)

# Set from settings at startup (OCR_ENABLED); module-level so the loaders stay plain functions.
OCR_ENABLED = True


class UnsupportedFileError(ValueError):
    """The file type isn't supported, or nothing readable could be extracted from it."""


@dataclass(frozen=True)
class Section:
    text: str
    metadata: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class LoadedFile:
    sections: list[Section]
    tabular: bool = False  # spreadsheets and CSVs: the table chunker is the right default


def markdown_table(rows: Iterable[Iterable[object]]) -> str:
    """Render rows as a markdown table; the first non-empty row is the header."""
    cleaned = [[_cell(c) for c in row] for row in rows]
    cleaned = [row for row in cleaned if any(row)]
    if not cleaned:
        return ""
    width = max(len(row) for row in cleaned)
    padded = [row + [""] * (width - len(row)) for row in cleaned]
    lines = ["| " + " | ".join(row) + " |" for row in padded]
    lines.insert(1, "| " + " | ".join(["---"] * width) + " |")
    return "\n".join(lines)


def _cell(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split()).replace("|", "\\|")


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _load_text(data: bytes) -> LoadedFile:
    return LoadedFile([Section(_decode(data))])


def _load_delimited(data: bytes) -> LoadedFile:
    # Kept as raw CSV/TSV: the table chunker parses it by file extension.
    return LoadedFile([Section(_decode(data))], tabular=True)


# A page with less text than this is treated as scanned and read with OCR (if it has images).
MIN_TEXT_CHARS_PER_PAGE = 20


def _load_pdf(data: bytes) -> LoadedFile:
    """Text layer per page; pages without one (scans, photos of documents) are OCR'd from their images."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    sections: list[Section] = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if len(text) < MIN_TEXT_CHARS_PER_PAGE and (scanned := _ocr_page_images(page)):
            sections.append(Section(scanned, {"page": str(number), "ocr": "true"}))
        elif text:
            sections.append(Section(text, {"page": str(number)}))
    if not sections:
        raise UnsupportedFileError("no text found, even with OCR (blank or unreadable scan?)")
    return LoadedFile(sections)


def _ocr_page_images(page: object) -> str:
    if not OCR_ENABLED:
        return ""
    texts = []
    for image in getattr(page, "images", []):
        try:
            texts.append(ocr.default_engine().read(image.data))
        except Exception:
            logger.warning("OCR failed for a PDF page image", exc_info=True)
    return "\n".join(t for t in texts if t.strip()).strip()


def _load_image(data: bytes) -> LoadedFile:
    """Photos and scans (PNG, JPG, TIFF, BMP, WEBP) are read with OCR."""
    if not OCR_ENABLED:
        raise UnsupportedFileError("images need OCR, which is disabled (OCR_ENABLED=false)")
    text = ocr.default_engine().read(data)
    if not text.strip():
        raise UnsupportedFileError("no text found in the image")
    return LoadedFile([Section(text, {"ocr": "true"})])


def _load_xlsx(data: bytes) -> LoadedFile:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)  # data_only: values, not formulas
    try:
        sections = [
            Section(table, {"sheet": sheet.title})
            for sheet in workbook.worksheets
            if (table := markdown_table(sheet.iter_rows(values_only=True)))
        ]
    finally:
        workbook.close()
    return LoadedFile(sections, tabular=True)


def _load_docx(data: bytes) -> LoadedFile:
    from docx import Document as WordDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = WordDocument(io.BytesIO(data))
    parts: list[str] = []
    for block in document.iter_inner_content():  # paragraphs and tables, in document order
        if isinstance(block, Paragraph) and (text := block.text.strip()):
            is_heading = (block.style.name if block.style else "").startswith("Heading")
            parts.append(f"## {text}" if is_heading else text)
        elif isinstance(block, Table) and (table := markdown_table([c.text for c in r.cells] for r in block.rows)):
            parts.append(table)
    return LoadedFile([Section("\n\n".join(parts))])


def _load_pptx(data: bytes) -> LoadedFile:
    from pptx import Presentation

    sections: list[Section] = []
    for number, slide in enumerate(Presentation(io.BytesIO(data)).slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and (text := shape.text_frame.text.strip()):
                parts.append(text)
            elif shape.has_table:
                parts.append(markdown_table([c.text for c in r.cells] for r in shape.table.rows))
        if slide.has_notes_slide and (notes := slide.notes_slide.notes_text_frame.text.strip()):
            parts.append(f"Speaker notes: {notes}")
        if text := "\n\n".join(p for p in parts if p):
            sections.append(Section(text, {"slide": str(number)}))
    return LoadedFile(sections)


LOADERS: dict[str, Callable[[bytes], LoadedFile]] = {
    ".txt": _load_text,
    ".md": _load_text,
    ".markdown": _load_text,
    ".json": _load_text,
    ".csv": _load_delimited,
    ".tsv": _load_delimited,
    ".pdf": _load_pdf,
    ".xlsx": _load_xlsx,
    ".xlsm": _load_xlsx,
    ".docx": _load_docx,
    ".pptx": _load_pptx,
    ".png": _load_image,
    ".jpg": _load_image,
    ".jpeg": _load_image,
    ".tif": _load_image,
    ".tiff": _load_image,
    ".bmp": _load_image,
    ".webp": _load_image,
}

SUPPORTED_EXTENSIONS = tuple(LOADERS)
_LEGACY_FORMATS = {".xls": ".xlsx", ".doc": ".docx", ".ppt": ".pptx"}


def load_file(filename: str, data: bytes) -> LoadedFile:
    extension = PurePath(filename).suffix.lower()
    if extension in _LEGACY_FORMATS:
        raise UnsupportedFileError(f"old {extension} format: save it as {_LEGACY_FORMATS[extension]} and upload again")
    loader = LOADERS.get(extension)
    if loader is None:
        raise UnsupportedFileError(f"unsupported file type {extension or '(none)'}")
    try:
        loaded = loader(data)
    except UnsupportedFileError:
        raise
    except Exception as exc:  # corrupt or password-protected files raise library-specific errors
        logger.warning("could not parse upload", extra={"upload": filename}, exc_info=True)
        raise UnsupportedFileError(f"could not read the file ({type(exc).__name__})") from exc
    if not any(s.text.strip() for s in loaded.sections):
        raise UnsupportedFileError("the file contains no text")
    return loaded
