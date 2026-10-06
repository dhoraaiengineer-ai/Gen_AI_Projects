"""Build small real PDF / Excel / Word / PowerPoint files in memory for loader and upload tests."""

import io


def make_pdf(pages: list[str]) -> bytes:
    """Minimal valid PDF with one line of Helvetica text per page (no PDF writer dependency needed)."""
    page_ids = [4 + 2 * i for i in range(len(pages))]
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] /Count {len(pages)} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for page_id, text in zip(page_ids, pages, strict=True):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_id + 1} 0 R >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")

    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return out


def make_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_docx(heading: str, paragraph: str, table: list[list[str]]) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(paragraph)
    word_table = document.add_table(rows=len(table), cols=len(table[0]))
    for r, row in enumerate(table):
        for c, value in enumerate(row):
            word_table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx(title: str, body: str, notes: str, table: list[list[str]]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])  # title + content
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    slide.notes_slide.notes_text_frame.text = notes

    second = presentation.slides.add_slide(presentation.slide_layouts[5])  # title only
    second.shapes.title.text = "Numbers"
    shape = second.shapes.add_table(len(table), len(table[0]), Inches(1), Inches(2), Inches(6), Inches(2))
    for r, row in enumerate(table):
        for c, value in enumerate(row):
            shape.table.cell(r, c).text = value
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
