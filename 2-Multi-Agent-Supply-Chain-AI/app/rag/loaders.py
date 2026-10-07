"""Document loaders. Parse uploads in memory into ordered sections with citation locations.

| Type        | Section unit                                  |
|-------------|-----------------------------------------------|
| PDF         | one per page (OCR when a page has no text)    |
| XLSX        | one markdown table per sheet (values)         |
| DOCX        | headings, paragraphs and tables in order      |
| PPTX        | one per slide, with tables and speaker notes  |
| CSV/TSV     | one markdown table                            |
| TXT/MD/JSON | text                                          |
| PNG/JPG/TIF | OCR                                           |
"""

from __future__ import annotations

import csv
import io
import json
import logging
from dataclasses import dataclass

from app.core.ports import OcrEngine

logger = logging.getLogger(__name__)

SUPPORTED = {"pdf", "docx", "xlsx", "pptx", "csv", "tsv", "txt", "md", "json", "png", "jpg", "jpeg", "tif", "tiff"}
LEGACY = {"xls": "Excel 97-2003 (.xls)", "doc": "Word 97-2003 (.doc)", "ppt": "PowerPoint 97-2003 (.ppt)"}
MIN_PDF_PAGE_CHARS = 25


class UnsupportedDocument(ValueError):
    pass


@dataclass
class Section:
    text: str
    location: str  # e.g. "p. 3", "sheet Rates", "slide 4"
    ocr: bool = False
    confidence: float | None = None
    is_table: bool = False


def extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def normalise(text: str) -> str:
    """CRLF → LF before hashing and chunking (Windows line endings would make identical files look changed)."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [line.rstrip() for line in text.split("\n")]
    out, blank = [], 0
    for line in lines:
        blank = blank + 1 if not line else 0
        if blank <= 1:
            out.append(line)
    return "\n".join(out).strip()


def _md_table(rows: list[list[str]]) -> str:
    rows = [[(c or "").replace("|", "/").replace("\n", " ").strip() for c in r] for r in rows if any((c or "").strip() for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, body = rows[0], rows[1:]
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def _ocr(engine: OcrEngine | None, image: bytes, location: str) -> Section:
    if engine is None:
        raise UnsupportedDocument("This file is a scanned image and OCR is disabled. Set OCR_ENGINE=tesseract or vision_llm.")
    result = engine.image_to_text(image)
    return Section(normalise(result.text), location, ocr=True, confidence=result.confidence)


def load_pdf(data: bytes, ocr: OcrEngine | None) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise UnsupportedDocument("This PDF is password-protected. Remove the password and upload again.")
    sections = []
    for i, page in enumerate(reader.pages, start=1):
        text = normalise(page.extract_text() or "")
        if len(text) >= MIN_PDF_PAGE_CHARS:
            sections.append(Section(text, f"p. {i}"))
            continue
        # No text layer → scanned page: render with pdfium and OCR it.
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(data)
        bitmap = pdf[i - 1].render(scale=300 / 72).to_pil()
        buf = io.BytesIO()
        bitmap.save(buf, format="PNG")
        sections.append(_ocr(ocr, buf.getvalue(), f"p. {i}"))
    return sections


def load_docx(data: bytes) -> list[Section]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    sections: list[Section] = []
    heading, buffer = "Document", []

    def flush() -> None:
        if buffer:
            sections.append(Section(normalise("\n".join(buffer)), heading[:100]))
            buffer.clear()

    for block in document.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(block, document)
            text = para.text.strip()
            if not text:
                continue
            if para.style is not None and para.style.name.lower().startswith("heading"):
                flush()
                heading = text
                buffer.append(f"## {text}")
            else:
                buffer.append(text)
        elif tag == "tbl":
            table = Table(block, document)
            buffer.append(_md_table([[cell.text for cell in row.cells] for row in table.rows]))
    flush()
    return sections


def load_xlsx(data: bytes) -> list[Section]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)  # values, not formulas
    sections = []
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        table = _md_table(rows)
        if table:
            sections.append(Section(table, f"sheet {ws.title}", is_table=True))
    return sections


def load_pptx(data: bytes) -> list[Section]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    sections = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                parts.append(_md_table([[c.text for c in r.cells] for r in shape.table.rows]))
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip())
        if parts:
            sections.append(Section(normalise("\n".join(parts)), f"slide {i}"))
    return sections


def load_delimited(data: bytes, delimiter: str) -> list[Section]:
    text = data.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(normalise(text)), delimiter=delimiter))
    table = _md_table(rows)
    return [Section(table, "table", is_table=True)] if table else []


def load_text(data: bytes, ext: str) -> list[Section]:
    text = data.decode("utf-8-sig", errors="replace")
    if ext == "json":
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except json.JSONDecodeError as exc:
            raise UnsupportedDocument(f"Invalid JSON: {exc.msg} (line {exc.lineno}).") from exc
    text = normalise(text)
    return [Section(text, "text")] if text else []


def load(filename: str, data: bytes, ocr: OcrEngine | None = None) -> list[Section]:
    ext = extension(filename)
    if ext in LEGACY:
        raise UnsupportedDocument(f"{LEGACY[ext]} files aren't supported. Save as .{ext}x and upload again.")
    if ext not in SUPPORTED:
        raise UnsupportedDocument(f".{ext or '?'} files aren't supported. Upload PDF, Word, Excel, PowerPoint, CSV, text or images.")
    if ext == "pdf":
        sections = load_pdf(data, ocr)
    elif ext == "docx":
        sections = load_docx(data)
    elif ext == "xlsx":
        sections = load_xlsx(data)
    elif ext == "pptx":
        sections = load_pptx(data)
    elif ext in {"csv", "tsv"}:
        sections = load_delimited(data, "\t" if ext == "tsv" else ",")
    elif ext in {"png", "jpg", "jpeg", "tif", "tiff"}:
        sections = [_ocr(ocr, data, "image")]
    else:
        sections = load_text(data, ext)
    sections = [s for s in sections if s.text.strip()]
    if not sections:
        raise UnsupportedDocument("No readable text was found in this file.")
    return sections
