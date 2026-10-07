"""Generate the fictional sample knowledge base in every supported format, plus a golden Q&A set.

    uv run python scripts/generate_sample_docs.py      # writes data/sample_docs/ and evals/golden/rag_golden.json

Every golden evidence quote must appear verbatim in its document (checked by tests/rag/test_sample_docs.py).
All organisations, people and terms are fictional.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from fpdf import FPDF
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "sample_docs"
GOLDEN = ROOT / "evals" / "golden" / "rag_golden.json"

PROCUREMENT = {
    "title": "Northwind Industrial - Global Procurement Policy v4.2",
    "pages": [
        (
            "1. Purpose and Scope",
            [
                "This policy governs the purchase of direct and indirect materials by all Northwind Industrial business units.",
                "Effective date: 2026-01-01. It applies to every purchase order issued from an approved ERP system.",
            ],
        ),
        (
            "4. Approval Thresholds",
            [
                "4.1 Purchase orders up to USD 10,000 may be released by a Buyer.",
                "4.2 Purchase orders above USD 50,000 require approval from a Procurement Manager or above before release to the supplier.",
                "4.3 Purchase orders above USD 250,000 additionally require approval from the Chief Financial Officer.",
                "4.4 AI-generated purchase recommendations are drafts. They must never be released without the approvals in this section.",
            ],
        ),
        (
            "5. Expedited Orders",
            [
                "5.1 Expedited orders are permitted when projected days of cover fall below the supplier lead time.",
                "5.2 The expedite premium must not exceed 12% of the standard unit price.",
                "5.3 Air freight for expedited orders requires Logistics Manager sign-off when freight cost exceeds USD 5,000.",
            ],
        ),
        (
            "6. Supplier Diversification",
            [
                "6.1 No critical SKU may rely on a single supplier for more than 70% of annual volume.",
                "6.2 A qualified secondary supplier must be maintained for all Power Systems components.",
            ],
        ),
    ],
}

INVENTORY_MD = """# Safety Stock and Service Level Standard

Effective date: 2026-02-01

## Service levels
- A-class items target a 95% cycle service level.
- B-class items target a 90% cycle service level.
- C-class items target an 85% cycle service level.

## Safety stock method
Safety stock is calculated as z multiplied by the standard deviation of daily demand multiplied by the square root of the lead time in days.
Demand variability is measured over the trailing 90 days.

## Review cadence
Reorder points are recalculated weekly. The planning review period is 30 days.
A stockout risk is critical when days of cover are below the supplier lead time and inbound stock does not restore the reorder point.
"""

CONTRACTS = [
    (
        "Nordvolt Energy - Master Supply Agreement.docx",
        "Nordvolt Energy",
        "SUP-003",
        "2026-01-01",
        "2028-06-30",
        [
            "Supplier: Nordvolt Energy (SUP-003). Buyer: Northwind Industrial.",
            "Products: lithium-ion battery packs including SKU-100 Lithium-Ion Battery Pack 48V.",
            "Clause 6.1 Lead time: Supplier guarantees a maximum lead time of three (3) business days for standard battery pack SKUs.",
            "Clause 6.2 Late delivery: Supplier grants a late-delivery credit of 2% of order value per day of delay, capped at 10%.",
            "Clause 7.1 Pricing: The unit price for SKU-100 is USD 11.00 for orders of 500 units or more.",
            "Clause 9.3 Quality: Defect rates above 0.5% trigger a corrective action plan within 10 business days.",
        ],
    ),
    (
        "Kaito Precision - Supply Agreement 2025.docx",
        "Kaito Precision Industries",
        "SUP-001",
        "2025-11-01",
        "2027-11-30",
        [
            "Supplier: Kaito Precision Industries (SUP-001). Buyer: Northwind Industrial.",
            "Clause 4.1 Lead time: Standard lead time is five (5) business days from purchase order acknowledgement.",
            "Clause 4.4 Minimum order quantity: The minimum order quantity for battery packs is 500 units.",
            "Clause 8.2 Payment terms: Net 45 days from invoice date.",
        ],
    ),
    (
        "Pacific Rim Components - Supply Contract 2023.docx",
        "Pacific Rim Components",
        "SUP-004",
        "2023-01-01",
        "2025-12-31",
        [
            "Supplier: Pacific Rim Components (SUP-004). Buyer: Northwind Industrial.",
            "Clause 3.2 Lead time: Standard lead time is fourteen (14) days.",
            "Clause 5.1 Pricing: The unit price for SKU-100 is USD 9.60.",
        ],
    ),
]

SHIPPING_SLIDES = [
    ("International Shipping Policy", ["Northwind Industrial Logistics", "Effective date: 2026-03-01"], "Owner: Global Logistics team."),
    ("Default Incoterms", ["Ocean freight from Asia uses FOB origin port.", "Road freight within the EU uses DAP."], ""),
    (
        "Lithium Battery Shipments",
        [
            "Lithium battery shipments by air must be declared as UN3480 dangerous goods.",
            "Batteries shipped by air are limited to 30% state of charge.",
        ],
        "Speaker notes: carriers reject undeclared UN3480 cargo; always attach the dangerous goods declaration.",
    ),
    (
        "Delay Escalation",
        [
            "Shipments delayed more than 3 days must be escalated to the Logistics Manager.",
            "Expedite by air only when the delay creates a critical stockout risk.",
        ],
        "",
    ),
]

RATE_CARD = [
    ["Lane", "Mode", "Rate USD per kg", "Transit days", "Fuel surcharge"],
    ["Shenzhen to Tokyo", "Ocean", 0.18, 21, "0%"],
    ["Shenzhen to Tokyo", "Air", 4.60, 3, "7%"],
    ["Busan to Osaka", "Ocean", 0.15, 6, "0%"],
    ["Hamburg to Rotterdam", "Road", 0.42, 2, "3%"],
    ["Gdansk to Rotterdam", "Rail", 0.27, 9, "2%"],
]

SCORECARD = [
    ["supplier_id", "supplier", "on_time_rate", "defect_rate", "fill_rate", "risk"],
    ["SUP-001", "Kaito Precision Industries", "94%", "0.6%", "97%", "low"],
    ["SUP-002", "Meridian Power Systems", "91%", "1.1%", "95%", "low"],
    ["SUP-003", "Nordvolt Energy", "98%", "0.3%", "99%", "low"],
    ["SUP-004", "Pacific Rim Components", "86%", "1.9%", "90%", "high"],
    ["SUP-005", "Atlas Industrial Supply", "95%", "0.5%", "96%", "low"],
]

INVOICE_LINES = [
    "FREIGHT INVOICE FI-20914",
    "Carrier: Bluewater Ocean Lines",
    "Shipment: SHP-48210",
    "Route: Shenzhen to Tokyo DC",
    "Weight: 4,250 kg",
    "Ocean freight: USD 765.00",
    "Port congestion surcharge: USD 180.00",
    "Total due: USD 945.00",
]


def write_pdf(path: Path) -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    for i, (heading, paragraphs) in enumerate(PROCUREMENT["pages"]):
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14)
        if i == 0:
            pdf.multi_cell(0, 9, PROCUREMENT["title"])
            pdf.ln(2)
        pdf.multi_cell(0, 8, heading)
        pdf.set_font("Helvetica", "", 11)
        for p in paragraphs:
            pdf.ln(2)
            pdf.multi_cell(0, 6, p)
    pdf.output(str(path))


def write_contract(path: Path, supplier: str, sid: str, effective: str, expiry: str, clauses: list[str]) -> None:
    doc = Document()
    doc.add_heading(f"{supplier} - Supply Agreement", level=1)
    doc.add_paragraph(f"Effective date: {effective}. Expiry date: {expiry}.")
    doc.add_heading("Terms", level=2)
    for c in clauses:
        doc.add_paragraph(c)
    table = doc.add_table(rows=2, cols=3)
    for j, h in enumerate(["Supplier ID", "Effective", "Expiry"]):
        table.rows[0].cells[j].text = h
    for j, v in enumerate([sid, effective, expiry]):
        table.rows[1].cells[j].text = v
    doc.save(str(path))


def write_pptx(path: Path) -> None:
    prs = Presentation()
    for title, bullets, notes in SHIPPING_SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    table_slide = prs.slides.add_slide(prs.slide_layouts[5])
    table_slide.shapes.title.text = "Carrier Selection"
    shape = table_slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(8), Inches(1.5))
    for r, row in enumerate([["Priority", "Preferred mode"], ["Critical stockout", "Air"], ["Standard replenishment", "Ocean"]]):
        for c, val in enumerate(row):
            shape.table.cell(r, c).text = val
            shape.table.cell(r, c).text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    prs.save(str(path))


def write_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rates"
    for row in RATE_CARD:
        ws.append(row)
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Note"])
    ws2.append(["Rates valid for Q4 shipments booked before 2026-12-15."])
    wb.save(str(path))


def write_csv(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(SCORECARD)


def write_scan(path: Path) -> None:
    img = Image.new("RGB", (1700, 1100), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 44)
    except OSError:
        font = ImageFont.load_default(size=44)
    y = 80
    for line in INVOICE_LINES:
        draw.text((90, y), line, fill="black", font=font)
        y += 100
    img.save(str(path), dpi=(300, 300))


GOLDEN_ITEMS = [
    {
        "question": "What approval is needed for a purchase order above USD 50,000?",
        "answer_contains": ["Procurement Manager"],
        "evidence": "Purchase orders above USD 50,000 require approval from a Procurement Manager or above before release to the supplier.",
        "source": "Global Procurement Policy v4.2.pdf",
    },
    {
        "question": "When are expedited orders permitted?",
        "answer_contains": ["lead time"],
        "evidence": "Expedited orders are permitted when projected days of cover fall below the supplier lead time.",
        "source": "Global Procurement Policy v4.2.pdf",
    },
    {
        "question": "What is the maximum expedite premium?",
        "answer_contains": ["12%"],
        "evidence": "The expedite premium must not exceed 12% of the standard unit price.",
        "source": "Global Procurement Policy v4.2.pdf",
    },
    {
        "question": "What share of annual volume may a single supplier have for a critical SKU?",
        "answer_contains": ["70%"],
        "evidence": "No critical SKU may rely on a single supplier for more than 70% of annual volume.",
        "source": "Global Procurement Policy v4.2.pdf",
    },
    {
        "question": "What lead time does Nordvolt guarantee for battery packs?",
        "answer_contains": ["three", "3"],
        "evidence": "Supplier guarantees a maximum lead time of three (3) business days for standard battery pack SKUs.",
        "source": "Nordvolt Energy - Master Supply Agreement.docx",
    },
    {
        "question": "What late-delivery credit does Nordvolt give?",
        "answer_contains": ["2%"],
        "evidence": "Supplier grants a late-delivery credit of 2% of order value per day of delay, capped at 10%.",
        "source": "Nordvolt Energy - Master Supply Agreement.docx",
    },
    {
        "question": "What are Kaito's payment terms?",
        "answer_contains": ["45"],
        "evidence": "Net 45 days from invoice date.",
        "source": "Kaito Precision - Supply Agreement 2025.docx",
    },
    {
        "question": "What service level do A-class items target?",
        "answer_contains": ["95%"],
        "evidence": "A-class items target a 95% cycle service level.",
        "source": "Safety Stock and Service Level Standard.md",
    },
    {
        "question": "How must lithium batteries be declared when shipped by air?",
        "answer_contains": ["UN3480"],
        "evidence": "Lithium battery shipments by air must be declared as UN3480 dangerous goods.",
        "source": "International Shipping Policy.pptx",
    },
    {
        "question": "When must a delayed shipment be escalated?",
        "answer_contains": ["3 days"],
        "evidence": "Shipments delayed more than 3 days must be escalated to the Logistics Manager.",
        "source": "International Shipping Policy.pptx",
    },
    {
        "question": "What is the air freight rate from Shenzhen to Tokyo?",
        "answer_contains": ["4.6"],
        "evidence": "| Shenzhen to Tokyo | Air | 4.6 |",
        "source": "Carrier Rate Card Q4.xlsx",
    },
    {
        "question": "What is Nordvolt's on-time rate in the supplier scorecard?",
        "answer_contains": ["98%"],
        "evidence": "| SUP-003 | Nordvolt Energy | 98% |",
        "source": "Supplier Scorecard.csv",
    },
    {
        "question": "What does our policy say about the CEO's favourite colour?",
        "answer_contains": [],
        "evidence": None,
        "source": None,
        "expect_refusal": True,
    },
]


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    write_pdf(OUT / "Global Procurement Policy v4.2.pdf")
    (OUT / "Safety Stock and Service Level Standard.md").write_text(INVENTORY_MD, encoding="utf-8")
    for name, supplier, sid, eff, exp, clauses in CONTRACTS:
        write_contract(OUT / name, supplier, sid, eff, exp, clauses)
    write_pptx(OUT / "International Shipping Policy.pptx")
    write_xlsx(OUT / "Carrier Rate Card Q4.xlsx")
    write_csv(OUT / "Supplier Scorecard.csv")
    write_scan(OUT / "Freight Invoice FI-20914 (scan).png")
    GOLDEN.parent.mkdir(parents=True, exist_ok=True)
    GOLDEN.write_text(json.dumps(GOLDEN_ITEMS, indent=2), encoding="utf-8")
    print(f"wrote {len(list(OUT.iterdir()))} documents to {OUT} and {len(GOLDEN_ITEMS)} golden items")


if __name__ == "__main__":
    main()
