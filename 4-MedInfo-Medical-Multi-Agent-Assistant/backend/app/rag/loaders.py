"""Document loaders: PDF (per page), DOCX, PPTX, XLSX, CSV/TSV, text, and images. Scanned PDF pages and
images go through OCR. Every section keeps its page / sheet / slide / heading for citations.
Everything works on in-memory bytes; nothing is written to disk."""

from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass
from pathlib import PurePath
from typing import Literal

from app.core.errors import AppError
from app.rag.interfaces import OcrEngine

logger = logging.getLogger(__name__)

TEXT_EXTS = {".txt", ".md"}
TABLE_TEXT_EXTS = {".csv", ".tsv"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg"}
SUPPORTED_EXTS = {".pdf", ".docx", ".pptx", ".xlsx", *TEXT_EXTS, *TABLE_TEXT_EXTS, *IMAGE_EXTS}
# A PDF page with less extractable text than this is treated as scanned and sent to OCR.
MIN_PAGE_TEXT_CHARS = 25


class UnsupportedFileTypeError(AppError):
    status_code = 415
    public_message = "This file type is not supported."


class DocumentParseError(AppError):
    status_code = 422
    public_message = "The document could not be read."


@dataclass(frozen=True)
class Section:
    text: str
    kind: Literal["text", "table"] = "text"
    page: int | None = None
    sheet: str | None = None
    slide: int | None = None
    heading: str | None = None
    ocr: bool = False
    header: str | None = None  # table header row, repeated in every table chunk
    flagged: bool = False  # guardrails neutralised instruction-like text in this section


def extension(filename: str) -> str:
    return PurePath(filename).suffix.lower()


def normalise_newlines(text: str) -> str:
    """CRLF / CR -> LF, so the same content hashes the same on Windows and Linux."""
    return text.replace("\r\n", "\n").replace("\r", "\n")


def normalise_bytes_for_hash(filename: str, data: bytes) -> bytes:
    if extension(filename) in TEXT_EXTS | TABLE_TEXT_EXTS:
        return data.replace(b"\r\n", b"\n").replace(b"\r", b"\n")
    return data


def _cell(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split()).replace("|", "/")


def table_rows_to_section(rows: list[list[object]], **loc: object) -> Section | None:
    cleaned = [[_cell(c) for c in r] for r in rows]
    cleaned = [r for r in cleaned if any(r)]
    if not cleaned:
        return None
    width = max(len(r) for r in cleaned)
    lines = ["| " + " | ".join(r + [""] * (width - len(r))) + " |" for r in cleaned]
    return Section(text="\n".join(lines), kind="table", header=lines[0], **loc)  # type: ignore[arg-type]


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# ---- per-format loaders -------------------------------------------------------------------------
def _load_pdf(data: bytes, ocr: OcrEngine | None) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentParseError("Encrypted PDFs are not supported.")
    sections: list[Section] = []
    for i, page in enumerate(reader.pages, 1):
        text = normalise_newlines(page.extract_text() or "").strip()
        used_ocr = False
        if len(text) < MIN_PAGE_TEXT_CHARS and ocr is not None:
            ocr_parts = [ocr.image_to_text(img.data) for img in page.images]
            ocr_text = "\n".join(t for t in ocr_parts if t.strip()).strip()
            if ocr_text:
                text, used_ocr = ocr_text, True
        if text:
            sections.append(Section(text=text, page=i, ocr=used_ocr))
    return sections


def _load_docx(data: bytes) -> list[Section]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    sections: list[Section] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            sections.append(Section(text="\n".join(buffer), heading=heading))
            buffer.clear()

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            if (para.style.name or "").lower().startswith(("heading", "title")):
                flush()
                heading = text
            else:
                buffer.append(text)
        elif tag == "tbl":
            flush()
            table = Table(child, doc)
            sec = table_rows_to_section([[c.text for c in row.cells] for row in table.rows], heading=heading)
            if sec:
                sections.append(sec)
    flush()
    return sections


def _load_pptx(data: bytes) -> list[Section]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    sections: list[Section] = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else None
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip() and shape != slide.shapes.title:
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in row.cells] for row in shape.table.rows]
                sec = table_rows_to_section(rows, slide=i, heading=title)
                if sec:
                    sections.append(sec)
        if texts:
            sections.append(Section(text="\n".join(texts), slide=i, heading=title))
    return sections


def _load_xlsx(data: bytes) -> list[Section]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections: list[Section] = []
    try:
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            sec = table_rows_to_section(rows, sheet=ws.title)
            if sec:
                sections.append(sec)
    finally:
        wb.close()
    return sections


def _load_delimited(data: bytes, delimiter: str) -> list[Section]:
    reader = csv.reader(io.StringIO(normalise_newlines(_decode(data))), delimiter=delimiter)
    sec = table_rows_to_section([list(r) for r in reader])
    return [sec] if sec else []


def _load_text(data: bytes) -> list[Section]:
    text = normalise_newlines(_decode(data)).strip()
    return [Section(text=text)] if text else []


def _load_image(data: bytes, ocr: OcrEngine | None) -> list[Section]:
    if ocr is None:
        raise DocumentParseError("OCR is not available for images.")
    text = ocr.image_to_text(data).strip()
    return [Section(text=text, page=1, ocr=True)] if text else []


def load_document(filename: str, data: bytes, ocr: OcrEngine | None = None) -> list[Section]:
    ext = extension(filename)
    if ext not in SUPPORTED_EXTS:
        raise UnsupportedFileTypeError(f"Unsupported file type {ext or '(none)'}.")
    try:
        if ext == ".pdf":
            sections = _load_pdf(data, ocr)
        elif ext == ".docx":
            sections = _load_docx(data)
        elif ext == ".pptx":
            sections = _load_pptx(data)
        elif ext == ".xlsx":
            sections = _load_xlsx(data)
        elif ext in TABLE_TEXT_EXTS:
            sections = _load_delimited(data, "\t" if ext == ".tsv" else ",")
        elif ext in IMAGE_EXTS:
            sections = _load_image(data, ocr)
        else:
            sections = _load_text(data)
    except AppError:
        raise
    except Exception as exc:
        logger.warning("document parse failed", extra={"ext": ext, "error": type(exc).__name__})
        raise DocumentParseError() from exc
    logger.info(
        "document loaded",
        extra={"ext": ext, "sections": len(sections), "ocr_sections": sum(s.ocr for s in sections)},
    )
    return sections
