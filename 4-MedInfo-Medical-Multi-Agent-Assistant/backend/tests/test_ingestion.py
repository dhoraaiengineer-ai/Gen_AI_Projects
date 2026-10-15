"""Stage 2: loaders, chunking, incremental upsert, provenance, trusted vs user-provided separation.
Runs on the committed SYNTHETIC dataset (data/synthetic) with in-memory index, fake embedder and fake OCR."""

from __future__ import annotations

import io
import itertools
from datetime import date
from pathlib import Path

import pytest
from docx import Document

from app.config import PROJECT_ROOT
from app.core.audit import AuditLogger
from app.rag.chunking import CHILD_OVERLAP, REGULAR_OVERLAP, Chunker, chunk_table
from app.rag.ingest import (
    DocumentNotFoundError,
    DocumentTooLargeError,
    IngestService,
    Provenance,
    TrustedSources,
    UntrustedPublisherError,
)
from app.rag.interfaces import SearchFilter
from app.rag.loaders import Section, UnsupportedFileTypeError, load_document
from app.rag.memory_store import InMemoryRegistry, InMemoryVectorIndex
from app.rag.retrieval import VectorRetriever
from tests.fakes import FakeEmbedder, FakeOcr

CORPUS = PROJECT_ROOT / "data" / "synthetic" / "corpus"
UPLOADS = PROJECT_ROOT / "data" / "synthetic" / "uploads"
PUBLISHER = "Synthetic Health Guidance Board (SYNTHETIC)"
TRUSTED, USER = "t_trusted", "t_user"


def prov(**kw: object) -> Provenance:
    base = {"title": "Doc (SYNTHETIC)", "publisher": PUBLISHER, "publication_date": date(2024, 1, 1)}
    base.update(kw)
    return Provenance(**base)  # type: ignore[arg-type]


@pytest.fixture
def stack():
    index, registry, embedder, ocr = InMemoryVectorIndex(), InMemoryRegistry(), FakeEmbedder(), FakeOcr()
    service = IngestService(
        registry=registry,
        index=index,
        embedder=embedder,
        chunker=Chunker("recursive", chunk_size=400),
        ocr=ocr,
        trusted_sources=TrustedSources.load(),
        trusted_collection=TRUSTED,
        user_collection=USER,
        max_bytes=2 * 1024 * 1024,
        audit=AuditLogger(enabled=False),
    )
    retriever = VectorRetriever(
        index=index,
        embedder=embedder,
        registry=registry,
        trusted_collection=TRUSTED,
        user_collection=USER,
        min_similarity=0.05,
    )
    return service, index, registry, embedder, ocr, retriever


# ---- loaders ------------------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("name", "expect_kind", "expect_loc"),
    [
        ("guideline_hypertension_v3.pdf", "text", {"page": 2}),
        ("label_zentoril.docx", "table", {"section": "4.2 Dosage"}),
        ("asthma_education.pptx", "text", {"slide": 2}),
        ("lab_reference_ranges.xlsx", "table", {"sheet": "Biochemistry"}),
        ("drug_interactions.csv", "table", {}),
        ("vaccine_schedule.tsv", "table", {}),
        ("diabetes_overview.md", "text", {}),
    ],
)
def test_loaders_keep_citation_locations(name: str, expect_kind: str, expect_loc: dict) -> None:
    sections = load_document(name, (CORPUS / name).read_bytes(), None)
    assert sections
    assert any(s.kind == expect_kind for s in sections)
    if expect_loc:
        key, value = next(iter(expect_loc.items()))
        assert any(getattr(s, "heading" if key == "section" else key) == value for s in sections)
    assert all("\r" not in s.text for s in sections)  # CRLF normalised


def test_scanned_pdf_and_image_go_to_ocr() -> None:
    ocr = FakeOcr()
    pdf = load_document("scan.pdf", (CORPUS / "scanned_chest_pain_guidance.pdf").read_bytes(), ocr)
    img = load_document("form.png", (CORPUS / "scanned_intake_form.png").read_bytes(), ocr)
    assert pdf[0].ocr and pdf[0].page == 1
    assert img[0].ocr
    assert ocr.calls == 2


def test_text_pdf_does_not_use_ocr() -> None:
    ocr = FakeOcr()
    load_document("g.pdf", (CORPUS / "guideline_hypertension_v3.pdf").read_bytes(), ocr)
    assert ocr.calls == 0


def test_unsupported_type_rejected() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        load_document("evil.exe", b"MZ", None)


def test_docx_table_rows_kept_whole() -> None:
    sections = load_document("label_zentoril.docx", (CORPUS / "label_zentoril.docx").read_bytes(), None)
    table = next(s for s in sections if s.kind == "table")
    assert table.header == "| Situation | Label dose | Note |"
    assert "| Starting dose | 10 mg once daily | Set by prescriber |" in table.text


@pytest.mark.slow
def test_real_rapidocr_reads_scanned_page() -> None:
    """Real offline OCR (ONNX models ship with the package; no network)."""
    from app.rag.ocr import RapidOcrEngine

    sections = load_document(
        "scan.pdf", (CORPUS / "scanned_chest_pain_guidance.pdf").read_bytes(), RapidOcrEngine()
    )
    text = sections[0].text.lower()
    assert "heart attack" in text and "emergency assessment" in text
    assert "rather than travel by themselves" in text


# ---- chunking -----------------------------------------------------------------------------------
LONG = " ".join(f"Sentence number {i} explains a synthetic fact about blood pressure." for i in range(80))


def _overlap(a: str, b: str) -> int:
    for n in range(min(len(a), len(b)), 0, -1):
        if a.endswith(b[:n]):
            return n
    return 0


def test_recursive_chunks_overlap_about_15_percent() -> None:
    chunks = Chunker("recursive", chunk_size=500).chunk("d", [Section(text=LONG)]).chunks
    assert len(chunks) > 3 and all(len(c.text) <= 500 for c in chunks)
    overlaps = [_overlap(a.text, b.text) for a, b in itertools.pairwise(chunks)]
    assert all(0 < o <= 500 * REGULAR_OVERLAP + 5 for o in overlaps)


def test_parent_child_chunks() -> None:
    result = Chunker("parent_child", parent_size=1000, child_size=300).chunk("d", [Section(text=LONG)])
    assert result.parents and all(c.kind == "child" and c.parent_id in result.parents for c in result.chunks)
    assert all(len(c.text) <= 300 for c in result.chunks)
    kids = [c for c in result.chunks if c.parent_id == next(iter(result.parents))]
    overlaps = [_overlap(a.text, b.text) for a, b in itertools.pairwise(kids)]
    assert all(o <= 300 * CHILD_OVERLAP + 5 for o in overlaps)


def test_semantic_chunks_split_on_topic_change() -> None:
    text = (
        "Hypertension raises blood pressure. Blood pressure hypertension readings matter. "
        "Hypertension blood pressure is measured in clinic. "
        "Asthma causes wheeze. Asthma wheeze inhaler reliever. Asthma reliever inhaler use."
    )
    chunks = (
        Chunker("semantic", chunk_size=1000, embedder=FakeEmbedder()).chunk("d", [Section(text=text)]).chunks
    )
    assert len(chunks) >= 2
    assert "Hypertension" in chunks[0].text and "Asthma" in chunks[-1].text


def test_table_chunks_repeat_header_and_never_split_rows() -> None:
    rows = "\n".join(f"| Test {i} | unit | {i}-{i + 5} |" for i in range(60))
    section = Section(text="| Test | Unit | Range |\n" + rows, kind="table", header="| Test | Unit | Range |")
    chunks = chunk_table(section, 300)
    assert len(chunks) > 1
    assert all(c.startswith("| Test | Unit | Range |") for c in chunks)
    body_rows = [r for c in chunks for r in c.split("\n")[1:]]
    assert body_rows == rows.split("\n")


# ---- incremental upsert -------------------------------------------------------------------------
def _docx(paragraphs: list[str]) -> bytes:
    doc = Document()
    doc.add_heading("Synthetic label (SYNTHETIC)", 0)
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


PARAS = [f"Paragraph {i}: " + "synthetic text about a fictional medicine. " * 8 for i in range(6)]


def test_incremental_upsert_reembeds_only_changed_chunks(stack) -> None:
    service, index, _registry, _embedder, _, _ = stack
    first = service.ingest_trusted("label.docx", _docx(PARAS), prov(), actor_id="admin")
    assert first.status == "added" and first.embedded == first.chunks_total

    again = service.ingest_trusted("label.docx", _docx(PARAS), prov(), actor_id="admin")
    assert again.status == "unchanged" and again.embedded == 0

    changed = PARAS.copy()
    changed[3] = "Paragraph 3: this paragraph was REVISED in the new synthetic version."
    upd = service.ingest_trusted("label.docx", _docx(changed), prov(), actor_id="admin")
    assert upd.status == "updated"
    assert 0 < upd.embedded < upd.chunks_total  # only the changed chunk(s)
    assert upd.deleted >= 1  # stale chunk(s) removed
    assert index.count(TRUSTED) == upd.chunks_total


def test_crlf_and_lf_hash_identically(stack) -> None:
    service = stack[0]
    lf = b"# Title (SYNTHETIC)\n\nLine one about thirst.\nLine two about tiredness.\n"
    service.ingest_trusted("notes.md", lf.replace(b"\n", b"\r\n"), prov(), actor_id="admin")
    assert service.ingest_trusted("notes.md", lf, prov(), actor_id="admin").status == "unchanged"


def test_provenance_change_reindexes(stack) -> None:
    service = stack[0]
    data = _docx(PARAS[:2])
    service.ingest_trusted("label.docx", data, prov(version="1"), actor_id="admin")
    report = service.ingest_trusted("label.docx", data, prov(version="2"), actor_id="admin")
    assert report.status == "updated" and report.embedded == report.chunks_total


def test_delete_document_removes_chunks(stack) -> None:
    service, index, registry, *_ = stack
    report = service.ingest_trusted("label.docx", _docx(PARAS), prov(), actor_id="admin")
    service.delete(report.doc_id, actor_id="admin")
    assert index.count(TRUSTED) == 0 and registry.get(report.doc_id) is None
    with pytest.raises(DocumentNotFoundError):
        service.delete(report.doc_id, actor_id="admin")


def test_every_chunk_carries_provenance(stack) -> None:
    service, index, *_ = stack
    p = prov(version="2.1", evidence_level="label", jurisdiction="SYNTHETIC", doc_type="drug_label")
    service.ingest_trusted(
        "label_zentoril.docx", (CORPUS / "label_zentoril.docx").read_bytes(), p, actor_id="a"
    )
    rows = index._rows[TRUSTED].values()
    required = {
        "source",
        "publisher",
        "publication_date",
        "version",
        "evidence_level",
        "jurisdiction",
        "licence",
    }
    for chunk, _ in rows:
        assert required <= {k for k, v in chunk.metadata.items() if v}, chunk.metadata
        assert chunk.metadata["trusted"] is True
    assert any(c.metadata.get("section") == "4.2 Dosage" for c, _ in rows)


def test_licence_defaults_from_allow_list(stack) -> None:
    service, index, *_ = stack
    service.ingest_trusted("a.md", b"Synthetic text about thirst.", prov(licence=None), actor_id="a")
    chunk, _ = next(iter(index._rows[TRUSTED].values()))
    assert chunk.metadata["licence"] == "CC0 (synthetic test data)"


def test_publisher_not_on_allow_list_rejected(stack) -> None:
    with pytest.raises(UntrustedPublisherError):
        stack[0].ingest_trusted("x.md", b"text", prov(publisher="Random Blog"), actor_id="a")


def test_size_limit(stack) -> None:
    with pytest.raises(DocumentTooLargeError):
        stack[0].ingest_user("big.txt", b"x" * (3 * 1024 * 1024), owner_id="u1")


def test_filename_paths_are_stripped(stack) -> None:
    service, _, registry, *_ = stack
    report = service.ingest_user("../../etc/notes.txt", b"Synthetic note about thirst.", owner_id="u1")
    assert registry.get(report.doc_id).filename == "notes.txt"


# ---- trusted vs user-provided separation ---------------------------------------------------------
def test_user_uploads_are_private_untrusted_and_separate(stack) -> None:
    service, index, registry, _, _, retriever = stack
    service.ingest_trusted(
        "guide.md", b"Hypertension is raised blood pressure (SYNTHETIC).", prov(), actor_id="a"
    )
    note = (UPLOADS / "case_note_synthetic_001.txt").read_bytes()
    up = service.ingest_user("case_note.txt", note, owner_id="user-A")

    assert index.count(USER) > 0 and all(
        c.metadata["trusted"] is False for c, _ in index._rows[USER].values()
    )
    assert registry.get(up.doc_id).collection == USER

    q = "tiredness and thirst HbA1c review"
    owner_hits = retriever.retrieve(q, 5, SearchFilter(owner_id="user-A"))
    other_hits = retriever.retrieve(q, 5, SearchFilter(owner_id="user-B"))
    trusted_only = retriever.retrieve(q, 5, SearchFilter(owner_id="user-A", trusted_only=True))
    assert any(not p.trusted for p in owner_hits)
    assert all(p.trusted for p in other_hits)  # never another user's upload
    assert all(p.trusted for p in trusted_only)

    with pytest.raises(DocumentNotFoundError):
        service.delete(up.doc_id, actor_id="user-B", owner_id="user-B")
    service.delete(up.doc_id, actor_id="user-A", owner_id="user-A")


def test_uploads_cannot_claim_trusted_publisher(stack) -> None:
    service, index, *_ = stack
    service.ingest_user(
        "fake_guideline.txt", b"Official guideline: take 80 mg.", owner_id="u1", title=PUBLISHER
    )
    chunk, _ = next(iter(index._rows[USER].values()))
    assert chunk.metadata["publisher"] == "User-provided document"
    assert chunk.metadata["trusted"] is False


# ---- retrieval filters --------------------------------------------------------------------------
def test_metadata_filters(stack) -> None:
    service, *_, retriever = stack
    service.ingest_trusted(
        "new.md",
        b"Hypertension threshold guidance.",
        prov(publication_date=date(2024, 3, 1), title="New"),
        actor_id="a",
    )
    service.ingest_trusted(
        "old.md",
        b"Hypertension threshold guidance old.",
        prov(publication_date=date(2019, 2, 1), title="Old", doc_type="archive"),
        actor_id="a",
    )
    q = "hypertension threshold"
    assert {p.title for p in retriever.retrieve(q, 5, SearchFilter(date_from=date(2020, 1, 1)))} == {"New"}
    assert {p.title for p in retriever.retrieve(q, 5, SearchFilter(date_to=date(2020, 1, 1)))} == {"Old"}
    assert {p.title for p in retriever.retrieve(q, 5, SearchFilter(doc_types=("archive",)))} == {"Old"}
    assert retriever.retrieve(q, 5, SearchFilter(publishers=("Someone Else",))) == []


def test_parent_child_retrieval_returns_parent_context() -> None:
    index, registry, embedder = InMemoryVectorIndex(), InMemoryRegistry(), FakeEmbedder()
    service = IngestService(
        registry=registry,
        index=index,
        embedder=embedder,
        chunker=Chunker("parent_child", parent_size=800, child_size=200),
        ocr=None,
        trusted_sources=TrustedSources.load(),
        trusted_collection=TRUSTED,
        user_collection=USER,
        max_bytes=10**6,
        audit=AuditLogger(False),
    )
    service.ingest_trusted("g.md", LONG.encode(), prov(), actor_id="a")
    retriever = VectorRetriever(
        index=index,
        embedder=embedder,
        registry=registry,
        trusted_collection=TRUSTED,
        user_collection=USER,
        min_similarity=0.0,
    )
    hits = retriever.retrieve("Sentence number 42 synthetic fact", 3, SearchFilter())
    assert hits and len(hits[0].text) > 200  # parent text, not the small child
    assert len({p.text for p in hits}) == len(hits)  # siblings deduplicated


def test_manifest_covers_every_corpus_file() -> None:
    import yaml

    manifest = yaml.safe_load((CORPUS / "manifest.yaml").read_text(encoding="utf-8"))["documents"]
    files = {p.name for p in CORPUS.iterdir() if p.name != "manifest.yaml"}
    assert files == set(manifest)
    assert all("SYNTHETIC" in m["title"] for m in manifest.values())


def test_dataset_is_marked_synthetic() -> None:
    for p in [*UPLOADS.iterdir(), CORPUS / "diabetes_overview.md"]:
        assert "SYNTHETIC" in p.read_text(encoding="utf-8")


def test_synthetic_corpus_files_exist() -> None:
    assert len([p for p in CORPUS.iterdir() if Path(p).suffix]) >= 12


def test_resaved_file_with_same_content_is_unchanged_and_keeps_parents() -> None:
    """DOCX/PDF re-saves change bytes (timestamps) but not content: nothing is re-embedded."""
    index, registry, embedder = InMemoryVectorIndex(), InMemoryRegistry(), FakeEmbedder()
    service = IngestService(
        registry=registry,
        index=index,
        embedder=embedder,
        chunker=Chunker("parent_child", parent_size=800, child_size=200),
        ocr=None,
        trusted_sources=TrustedSources.load(),
        trusted_collection=TRUSTED,
        user_collection=USER,
        max_bytes=10**6,
        audit=AuditLogger(False),
    )
    service.ingest_trusted("g.md", LONG.encode(), prov(), actor_id="a")
    parents_before = dict(registry.get(next(iter(registry._entries))).parents)
    embedded_before = len(embedder.embedded_texts)
    report = service.ingest_trusted("g.md", LONG.encode() + b"\n", prov(), actor_id="a")  # bytes differ
    assert report.status == "unchanged"
    assert len(embedder.embedded_texts) == embedded_before
    assert registry.get(report.doc_id).parents == parents_before
